"""Lightweight parser tests (docling replacement)."""
import io

from app.services.ingest import parse_file


def _docx_bytes():
    import docx

    d = docx.Document()
    d.add_heading("Refund Policy", level=2)
    d.add_paragraph("Items may be returned within thirty days of purchase, no questions asked.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Zone"
    t.rows[0].cells[1].text = "Days"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["zone", "price"])
    for zone, price in (("north", 5), ("south", 6), ("east", 7), ("west", 8), ("central", 9)):
        ws.append([zone, price])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes():
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Delivery SLA"
    body = slide.placeholders[1].text_frame
    body.text = "Zones A and B are delivered within two business days."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_parse_docx_headings_and_tables():
    chunks = parse_file(_docx_bytes(), "docx")
    text = "\n".join(c.page_content for c in chunks)
    assert "thirty days" in text and "Zone" in text
    assert chunks[0].metadata["section"] == "Refund Policy"


def test_parse_xlsx_sheets_become_sections():
    chunks = parse_file(_xlsx_bytes(), "xlsx")
    assert chunks and chunks[0].metadata["section"] == "Rates"
    assert "north | 5" in chunks[0].page_content and "central | 9" in chunks[0].page_content


def test_parse_pptx_titles_become_sections():
    chunks = parse_file(_pptx_bytes(), "pptx")
    text = "\n".join(c.page_content for c in chunks)
    assert "two business days" in text
    assert chunks[0].metadata["section"] == "Delivery SLA"


def test_parse_html_headings_and_text():
    body = ("We ship to every zone worldwide and dispatch all orders within one "
            "business day of payment confirmation.")
    html = (f"<html><head><style>.x{{}}</style></head><body>"
            f"<h2>Shipping</h2><p>{body}</p>"
            f"<script>evil()</script></body></html>").encode()
    chunks = parse_file(html, "html")
    text = "\n".join(c.page_content for c in chunks)
    assert "dispatch all orders" in text
    assert "evil()" not in text and ".x{}" not in text
    assert chunks[0].metadata["section"] == "Shipping"


def test_parse_blank_pdf_yields_no_chunks():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    assert parse_file(buf.getvalue(), "pdf") == []
