import asyncio
import io
import time
from datetime import datetime, timedelta, timezone
from html.parser import HTMLParser
from pathlib import Path

from langchain_core.documents import Document
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from app.config import Settings
from app.models import Document as DocRow
from app.models import Workspace
from app.rag.chunker import chunk_markdown, chunk_plain_text

SUPPORTED_EXTS = frozenset({"pdf", "docx", "pptx", "xlsx", "html", "md"})
_RETRY_DELAYS = [2, 4, 8]
_MAX_SHEET_ROWS = 2000


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(p.extract_text() or "" for p in reader.pages)


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower()
        lines.append(f"## {text}" if style.startswith(("heading", "title")) else text)
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n\n".join(lines)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for slide in prs.slides:
        title = slide.shapes.title.text.strip() if slide.shapes.title else ""
        if title:
            lines.append(f"## {title}")
        for shape in slide.shapes:
            if not shape.has_text_frame or shape is slide.shapes.title:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text and text != title:
                    lines.append(text)
        lines.append("")
    return "\n".join(lines)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"## {sheet.title}")
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= _MAX_SHEET_ROWS:
                lines.append(f"… truncated after {_MAX_SHEET_ROWS} rows")
                break
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
        lines.append("")
    return "\n".join(lines)


class _TextExtract(HTMLParser):
    _SKIP = {"script", "style", "noscript"}
    _HEADINGS = {"h1", "h2"}

    def __init__(self):
        super().__init__()
        self.chunks: list[str] = []
        self._skip_depth = 0
        self._heading = None
        self._buf: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._HEADINGS and self._heading is None:
            self._heading = tag

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._HEADINGS and self._heading == tag:
            text = " ".join("".join(self._buf).split())
            if text:
                self.chunks.append(f"## {text}")
            self._heading, self._buf = None, []

    def handle_data(self, data):
        if self._skip_depth:
            return
        if self._heading is not None:
            self._buf.append(data)
        else:
            text = " ".join(data.split())
            if text:
                self.chunks.append(text)


def _html_text(data: bytes) -> str:
    parser = _TextExtract()
    parser.feed(data.decode("utf-8", errors="replace"))
    return "\n\n".join(parser.chunks)


_PARSERS = {"pdf": _pdf_text, "docx": _docx_text, "pptx": _pptx_text,
            "xlsx": _xlsx_text, "html": _html_text}


def parse_file(data: bytes, ext: str) -> list[Document]:
    ext = ext.lower().lstrip(".")
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"Unsupported file type: .{ext}")
    if ext == "md":
        return chunk_markdown(data.decode("utf-8", errors="replace"))
    try:
        text = _PARSERS[ext](data)
    except ImportError as e:
        raise RuntimeError(f"parser dependency missing for .{ext}: {e}") from e
    except Exception as e:
        raise RuntimeError(f"parse failed: {e}") from e
    return chunk_markdown(text)


def create_document_row(
    session: Session, workspace: Workspace, filename: str, source: str,
    uploaded_by: str | None = None, byte_size: int = 0,
) -> DocRow:
    row = DocRow(
        workspace_id=workspace.id, filename=filename, file_ext=_ext(filename),
        byte_size=byte_size, source=source, uploaded_by=uploaded_by,
    )
    session.add(row)
    session.flush()
    return row


def find_inflight_duplicate(session: Session, workspace_id: str, filename: str,
                            byte_size: int, minutes: int = 10) -> DocRow | None:
    """Same file still being processed — Telegram redelivers webhooks on slow
    responses, so this guards against double-ingesting."""
    since = datetime.now(timezone.utc) - timedelta(minutes=minutes)
    return session.scalar(
        select(DocRow).where(
            DocRow.workspace_id == workspace_id, DocRow.filename == filename,
            DocRow.byte_size == byte_size, DocRow.status == "processing",
            DocRow.created_at >= since,
        ).order_by(DocRow.created_at.desc()).limit(1))


def save_original(settings: Settings, slug: str, document_id: str, ext: str, data: bytes) -> Path | None:
    import logging

    try:
        folder = Path(settings.upload_dir) / slug
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{document_id}.{ext}"
        path.write_bytes(data)
        return path
    except OSError as e:  # read-only FS (e.g. serverless) — originals are optional
        logging.getLogger(__name__).warning("could not save original file: %s", e)
        return None


def _add_with_retry(store, docs: list[Document]) -> None:
    last: Exception | None = None
    for delay in (*_RETRY_DELAYS, None):
        try:
            store.add_documents(docs)
            return
        except Exception as e:  # noqa: BLE001 — retry any store failure
            last = e
            if delay is not None:
                time.sleep(delay)
    raise last  # type: ignore[misc]


def ingest_document_sync(
    session_factory, document_id: str, data: bytes, ext: str, store,
    settings: Settings | None = None, _session: Session | None = None,
    source_file: str | None = None,
) -> int:
    if _session is not None:  # tests pass the live session directly
        row = _session.get(DocRow, document_id)
        try:
            chunks = parse_file(data, ext)
            for c in chunks:
                c.metadata["document_id"] = document_id
                if source_file:
                    c.metadata["source_file"] = source_file
            _add_with_retry(store, chunks)
            row.status, row.chunk_count, row.error = "ready", len(chunks), None
            _session.commit()
            return len(chunks)
        except Exception as e:
            _session.rollback()
            row = _session.get(DocRow, document_id)
            row.status, row.error = "failed", str(e)[:500]
            _session.commit()
            raise

    def _run(session: Session) -> int:
        row = session.get(DocRow, document_id)
        chunks = parse_file(data, ext)
        for c in chunks:
            c.metadata["document_id"] = document_id
            if source_file:
                c.metadata["source_file"] = source_file
        _add_with_retry(store, chunks)
        row.status, row.chunk_count, row.error = "ready", len(chunks), None
        session.commit()
        return len(chunks)

    try:
        with _open_session(session_factory) as session:
            return _run(session)
    except Exception as e:
        with _open_session(session_factory) as session:
            row = session.get(DocRow, document_id)
            row.status, row.error = "failed", str(e)[:500]
            session.commit()
        raise


async def ingest_document_async(session_factory, document_id: str, data: bytes, ext: str, store,
                                settings: Settings | None = None,
                                source_file: str | None = None) -> int:
    return await asyncio.to_thread(
        ingest_document_sync, session_factory, document_id, data, ext, store, settings,
        None, source_file)


def ingest_text_sync(
    session_factory, workspace: Workspace, title: str, text: str, store,
    source: str = "paste", uploaded_by: str | None = None,
    settings: Settings | None = None, _session: Session | None = None,
) -> DocRow:
    chunks = chunk_plain_text(text)
    if not chunks:
        raise ValueError("Content is too short to index.")
    if _session is not None:
        row = create_document_row(_session, workspace, title, source, uploaded_by, len(text.encode()))
        for c in chunks:
            c.metadata["document_id"] = row.id
            c.metadata["source_file"] = title
        _add_with_retry(store, chunks)
        row.status, row.chunk_count = "ready", len(chunks)
        _session.commit()
        return row

    with _open_session(session_factory) as session:
        row = create_document_row(session, workspace, title, source, uploaded_by, len(text.encode()))
        for c in chunks:
            c.metadata["document_id"] = row.id
            c.metadata["source_file"] = title
        _add_with_retry(store, chunks)
        row.status, row.chunk_count = "ready", len(chunks)
        session.commit()
        return row


class _open_session:
    def __init__(self, factory):
        self.factory = factory
        self.session = None

    def __enter__(self):
        self.session = self.factory()
        return self.session

    def __exit__(self, *exc):
        self.session.close()
        return False
